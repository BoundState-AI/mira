#!/usr/bin/env python3
"""
Module: MIUL Excel Generator
Purpose: Generate Materials Identification and Usage List (MIUL) in Excel format

BoundState AI 2026
Version: 2.0.0
Classification: BoundState Proprietary

Compliance: NASA-STD-8739.8B, SWE-013, SWE-033
Safety Critical: False
"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import PlanGeneratorSession
from ..policy import PolicyPack, create_default_policy_pack

LOGGER = logging.getLogger(__name__)

__version__ = "2.0.0"
__safety_critical__ = False


class MIULExcelGenerator:
    """
    Generate MIUL Excel spreadsheets.

    Produces standard NASA MIUL format with material identification,
    test data, applicability, and compliance status.
    """

    def __init__(
        self,
        session: PlanGeneratorSession,
        policy_pack: Optional[PolicyPack] = None,
    ):
        """
        Initialize generator.

        Args:
            session: Plan generator session.
            policy_pack: Optional policy pack.
        """
        self._session = session
        self._policy = policy_pack or create_default_policy_pack()

    def generate(self, output_path: Optional[Path] = None) -> bytes:
        """
        Generate MIUL Excel workbook.

        Args:
            output_path: Optional file path to save workbook.

        Returns:
            Workbook bytes.
        """
        wb = Workbook()

        # Remove default sheet
        default_sheet = wb.active
        wb.remove(default_sheet)

        # Create sheets
        self._create_summary_sheet(wb)
        self._create_materials_sheet(wb)
        self._create_environments_sheet(wb)
        self._create_applicability_sheet(wb)
        self._create_nonconformances_sheet(wb)
        self._create_closures_sheet(wb)

        # Save to bytes
        buffer = BytesIO()
        wb.save(buffer)
        xlsx_bytes = buffer.getvalue()

        # Optionally save to file
        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, "wb") as f:
                f.write(xlsx_bytes)
            LOGGER.info(f"Saved MIUL Excel to {output_path}")

        return xlsx_bytes

    def _create_summary_sheet(self, wb: Workbook) -> None:
        """Create summary sheet with comprehensive metrics."""
        ws = wb.create_sheet("Summary")

        # Header
        ws["A1"] = "Materials Identification and Usage List (MIUL)"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:D1")

        # Count materials by rating
        a_rated = b_rated = c_rated = not_resolved = 0
        for material in self._session.materials:
            resolution = self._session.get_resolution(material.id)
            if resolution:
                rating = resolution.flam_rating.upper() if resolution.flam_rating else "U"
                if rating == "A":
                    a_rated += 1
                elif rating == "B":
                    b_rated += 1
                elif rating in ("C", "D", "X"):
                    c_rated += 1
                else:
                    not_resolved += 1
            else:
                not_resolved += 1

        # Count test deltas outside envelope
        deltas_outside = 0
        for material in self._session.materials:
            resolution = self._session.get_resolution(material.id)
            if resolution and resolution.test_o2_pct:
                for env in self._session.environments:
                    if env.o2_pct > resolution.test_o2_pct:
                        deltas_outside += 1

        # Basic info section
        ws["A3"] = "DOCUMENT INFORMATION"
        ws["A3"].font = Font(bold=True, size=12)
        ws.merge_cells("A3:B3")

        basic_data = [
            ("Program", self._session.program_name or "TBD"),
            ("Phase", self._session.program_phase or "TBD"),
            ("Session ID", self._session.session_id),
            ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ]

        row = 4
        for label, value in basic_data:
            ws[f"A{row}"] = label
            ws[f"B{row}"] = str(value)
            ws[f"A{row}"].font = Font(bold=True)
            row += 1

        # Material counts section
        row += 1
        ws[f"A{row}"] = "MATERIAL SUMMARY"
        ws[f"A{row}"].font = Font(bold=True, size=12)
        ws.merge_cells(f"A{row}:B{row}")
        row += 1

        material_data = [
            ("Total Materials", len(self._session.materials)),
            ("A-Rated (Acceptable)", a_rated),
            ("B-Rated (Limited Use)", b_rated),
            ("C/D/X-Rated (Restricted)", c_rated),
            ("Not Resolved", not_resolved),
        ]

        for label, value in material_data:
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = Font(bold=True)
            # Highlight C/D/X and Not Resolved
            if "Restricted" in label and value > 0:
                ws[f"B{row}"].fill = PatternFill(start_color="FFCCCC", fill_type="solid")
            elif "Not Resolved" in label and value > 0:
                ws[f"B{row}"].fill = PatternFill(start_color="FFFFCC", fill_type="solid")
            row += 1

        # Compliance status section
        row += 1
        ws[f"A{row}"] = "COMPLIANCE STATUS"
        ws[f"A{row}"].font = Font(bold=True, size=12)
        ws.merge_cells(f"A{row}:B{row}")
        row += 1

        compliance_data = [
            ("Environments", len(self._session.environments)),
            ("Test Deltas Outside Envelope", deltas_outside),
            ("Nonconformances", len(self._session.nonconformances)),
            ("Closures Recorded", len(self._session.closure_decisions)),
            ("Open Gaps", len(self._session.accepted_open_gaps)),
        ]

        for label, value in compliance_data:
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = Font(bold=True)
            # Highlight issues
            if value > 0 and ("Outside" in label or "Nonconformances" in label or "Open" in label):
                ws[f"B{row}"].fill = PatternFill(start_color="FFCCCC", fill_type="solid")
            row += 1

        # Action items section
        row += 1
        ws[f"A{row}"] = "ACTION ITEMS"
        ws[f"A{row}"].font = Font(bold=True, size=12)
        ws.merge_cells(f"A{row}:B{row}")
        row += 1

        # Generate action items
        action_items = []
        if not_resolved > 0:
            action_items.append(f"Resolve MAPTIS codes for {not_resolved} material(s)")
        if deltas_outside > 0:
            action_items.append(f"Document rationale for {deltas_outside} test delta(s) outside envelope")
        if len(self._session.nonconformances) > len(self._session.closure_decisions) + len(self._session.accepted_open_gaps):
            open_nc = len(self._session.nonconformances) - len(self._session.closure_decisions) - len(self._session.accepted_open_gaps)
            action_items.append(f"Disposition {open_nc} open nonconformance(s)")

        if action_items:
            for i, item in enumerate(action_items, start=1):
                ws[f"A{row}"] = f"{i}."
                ws[f"B{row}"] = item
                ws[f"B{row}"].fill = PatternFill(start_color="FFFFCC", fill_type="solid")
                row += 1
        else:
            ws[f"A{row}"] = "None"
            ws[f"B{row}"] = "Document is ready for review"
            ws[f"B{row}"].fill = PatternFill(start_color="CCFFCC", fill_type="solid")

        # Adjust column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 50

    def _create_materials_sheet(self, wb: Workbook) -> None:
        """Create materials sheet."""
        ws = wb.create_sheet("Materials")

        # Headers
        headers = [
            "Item #",
            "Material Name",
            "Application",
            "MAPTIS Code",
            "Quantity",
            "Mass (g)",
            "Surface Area (cm²)",
            "FLAM Rating",
            "Test O2 %",
            "Status",
        ]
        self._add_header_row(ws, headers)

        # Data rows
        for i, material in enumerate(self._session.materials, start=2):
            resolution = self._session.get_resolution(material.id)

            # Check for nonconformances
            has_nc = any(
                nc.material_id == material.id for nc in self._session.nonconformances
            )
            status = "NC" if has_nc else "OK"

            row_data = [
                i - 1,
                material.name,
                material.application,
                material.maptis_code or "TBD",
                material.quantity,
                material.mass_g or "TBD",
                material.surface_area_cm2 or "TBD",
                resolution.flam_rating if resolution else "Not Resolved",
                resolution.test_o2_pct if resolution else "N/A",
                status,
            ]

            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=i, column=col, value=value)
                if status == "NC" and col == len(row_data):
                    cell.fill = PatternFill(start_color="FFCCCC", fill_type="solid")

        self._auto_fit_columns(ws, headers)

    def _create_environments_sheet(self, wb: Workbook) -> None:
        """Create environments sheet."""
        ws = wb.create_sheet("Environments")

        headers = [
            "Environment",
            "O2 %",
            "Pressure (psia)",
            "Temp Min (°C)",
            "Temp Max (°C)",
            "Habitable",
            "Vacuum",
            "Lunar Dust",
            "Phase",
        ]
        self._add_header_row(ws, headers)

        for i, env in enumerate(self._session.environments, start=2):
            row_data = [
                env.name,
                env.o2_pct,
                env.pressure_psia,
                env.temp_range_c[0],
                env.temp_range_c[1],
                "Yes" if env.habitable else "No",
                "Yes" if env.vacuum_exposure else "No",
                "Yes" if env.lunar_dust else "No",
                env.phase,
            ]
            for col, value in enumerate(row_data, start=1):
                ws.cell(row=i, column=col, value=value)

        self._auto_fit_columns(ws, headers)

    def _create_applicability_sheet(self, wb: Workbook) -> None:
        """Create applicability matrix sheet."""
        ws = wb.create_sheet("Applicability")

        if not self._session.applicability_matrix:
            ws["A1"] = "Applicability matrix not computed"
            return

        headers = ["Material/Environment", "MPR", "Status", "Rationale", "Triggered By"]
        self._add_header_row(ws, headers)

        row = 2
        for key, mprs in self._session.applicability_matrix.items():
            for mpr_id, decision in mprs.items():
                row_data = [
                    key,
                    mpr_id,
                    decision.get("status", "unknown"),
                    decision.get("rationale", ""),
                    decision.get("triggered_by", ""),
                ]
                for col, value in enumerate(row_data, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    # Color code by status
                    if col == 3:
                        if value == "applicable":
                            cell.fill = PatternFill(start_color="CCFFCC", fill_type="solid")
                        elif value == "not_applicable_by_condition":
                            cell.fill = PatternFill(start_color="FFFFCC", fill_type="solid")
                row += 1

        self._auto_fit_columns(ws, headers)

    def _create_nonconformances_sheet(self, wb: Workbook) -> None:
        """Create nonconformances sheet."""
        ws = wb.create_sheet("Nonconformances")

        headers = [
            "NC ID",
            "Material",
            "Environment",
            "Gap Type",
            "Failing Req",
            "Failing Value",
            "Required Value",
            "Severity",
            "Status",
        ]
        self._add_header_row(ws, headers)

        for i, nc in enumerate(self._session.nonconformances, start=2):
            material = self._session.get_material(nc.material_id)
            env = self._session.get_environment(nc.environment_id)

            # Determine status
            closure = next(
                (c for c in self._session.closure_decisions if c.gap_id == nc.nonconformance_id),
                None
            )
            accepted = next(
                (a for a in self._session.accepted_open_gaps if a.gap_id == nc.nonconformance_id),
                None
            )

            if closure:
                status = f"Closed: {closure.selected_path}"
            elif accepted:
                status = "Accepted Open"
            else:
                status = "Open"

            row_data = [
                nc.nonconformance_id[:8],
                material.name if material else nc.material_id,
                env.name if env else nc.environment_id,
                nc.gap_type.value.upper(),
                nc.failing_requirement,
                nc.failing_value,
                nc.required_value,
                nc.severity.value if nc.severity else "medium",
                status,
            ]

            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=i, column=col, value=value)
                # Color code by status
                if col == len(row_data):
                    if "Open" in str(value) and "Accepted" not in str(value):
                        cell.fill = PatternFill(start_color="FFCCCC", fill_type="solid")
                    elif "Closed" in str(value):
                        cell.fill = PatternFill(start_color="CCFFCC", fill_type="solid")
                    elif "Accepted" in str(value):
                        cell.fill = PatternFill(start_color="FFFFCC", fill_type="solid")

        self._auto_fit_columns(ws, headers)

    def _create_closures_sheet(self, wb: Workbook) -> None:
        """Create closures sheet."""
        ws = wb.create_sheet("Closures")

        headers = [
            "Gap ID",
            "Selected Path",
            "Rationale",
            "Decided By",
            "Decided At",
        ]
        self._add_header_row(ws, headers)

        for i, closure in enumerate(self._session.closure_decisions, start=2):
            row_data = [
                closure.gap_id[:8],
                closure.selected_path,
                closure.decision_rationale,
                closure.decided_by,
                closure.decided_at.strftime("%Y-%m-%d") if closure.decided_at else "N/A",
            ]
            for col, value in enumerate(row_data, start=1):
                ws.cell(row=i, column=col, value=value)

        self._auto_fit_columns(ws, headers)

    def _add_header_row(self, ws, headers: list[str]) -> None:
        """Add formatted header row."""
        header_fill = PatternFill(start_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

    def _auto_fit_columns(self, ws, headers: list[str]) -> None:
        """Auto-fit column widths."""
        for col, header in enumerate(headers, start=1):
            max_length = len(header)
            column_letter = get_column_letter(col)

            for row in range(2, ws.max_row + 1):
                cell_value = ws.cell(row=row, column=col).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))

            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


def generate_miul_xlsx(
    session: PlanGeneratorSession,
    output_path: Optional[Path] = None,
    policy_pack: Optional[PolicyPack] = None,
) -> bytes:
    """
    Generate MIUL Excel workbook.

    Args:
        session: Plan generator session.
        output_path: Optional file path to save workbook.
        policy_pack: Optional policy pack.

    Returns:
        Workbook bytes.
    """
    generator = MIULExcelGenerator(session, policy_pack)
    return generator.generate(output_path)
