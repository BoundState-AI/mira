#!/usr/bin/env python3
"""
Module: M&P Plan DOCX Generator
Purpose: Generate M&P Plan documents in DOCX format

BoundState AI 2026
Version: 2.0.0
Classification: BoundState Proprietary

Compliance: NASA-STD-8739.8B, SWE-013, SWE-033
Safety Critical: False
"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from ..models import PlanGeneratorSession
from ..policy import PolicyPack, create_default_policy_pack
from .section_builder import SectionBuilder, SectionContent

LOGGER = logging.getLogger(__name__)

__version__ = "2.0.0"
__safety_critical__ = False


class MPPlanDocxGenerator:
    """
    Generate M&P Plan DOCX documents.

    Uses pre-computed SectionContent - no business logic in this class.
    """

    def __init__(
        self,
        session: PlanGeneratorSession,
        policy_pack: Optional[PolicyPack] = None,
    ):
        """
        Initialize generator.

        Args:
            session: Plan generator session.
            policy_pack: Optional policy pack.
        """
        self._session = session
        self._policy = policy_pack or create_default_policy_pack()
        self._builder = SectionBuilder(session, policy_pack)

    def generate(self, output_path: Optional[Path] = None) -> bytes:
        """
        Generate M&P Plan DOCX.

        Args:
            output_path: Optional file path to save document.

        Returns:
            Document bytes.
        """
        doc = Document()

        # Set up styles
        self._setup_styles(doc)

        # Build sections
        sections = self._builder.build_all_sections()

        # Render each section
        for section in sections:
            self._render_section(doc, section)

        # Add footer with generation info
        self._add_footer(doc)

        # Save to bytes
        buffer = BytesIO()
        doc.save(buffer)
        doc_bytes = buffer.getvalue()

        # Optionally save to file
        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, "wb") as f:
                f.write(doc_bytes)
            LOGGER.info(f"Saved M&P Plan DOCX to {output_path}")

        return doc_bytes

    def _setup_styles(self, doc: Document) -> None:
        """Set up document styles."""
        styles = doc.styles

        # Heading 1
        if "MP Heading 1" not in [s.name for s in styles]:
            style = styles.add_style("MP Heading 1", WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Arial"
            style.font.size = Pt(16)
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)

        # Heading 2
        if "MP Heading 2" not in [s.name for s in styles]:
            style = styles.add_style("MP Heading 2", WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Arial"
            style.font.size = Pt(14)
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(4)

        # Normal text
        if "MP Normal" not in [s.name for s in styles]:
            style = styles.add_style("MP Normal", WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Arial"
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(6)

    def _render_section(self, doc: Document, section: SectionContent) -> None:
        """Render a section to the document."""
        if section.is_empty and section.section_id != "cover":
            return

        # Section heading
        if section.section_number:
            heading_text = f"{section.section_number} {section.section_title}"
        else:
            heading_text = section.section_title

        if section.section_id == "cover":
            # Cover page - centered, larger font
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(section.text_content)
            run.font.size = Pt(14)
            run.font.bold = True
            doc.add_page_break()
        else:
            # Regular section
            doc.add_heading(heading_text, level=1 if section.section_number.count(".") == 0 else 2)

        # Text content
        if section.text_content and section.section_id != "cover":
            doc.add_paragraph(section.text_content)

        # Bullet points
        if section.bullet_points:
            for point in section.bullet_points:
                doc.add_paragraph(point, style="List Bullet")

        # Table content
        if section.table_content:
            self._add_table(doc, section.table_content)

        # Subsections
        for subsection in section.subsections:
            self._render_section(doc, subsection)

    def _add_table(self, doc: Document, rows: list[dict]) -> None:
        """Add a table to the document."""
        if not rows:
            return

        # Get headers from first row
        headers = list(rows[0].keys())

        # Create table
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header row
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header.replace("_", " ").title()
            header_cells[i].paragraphs[0].runs[0].font.bold = True

        # Data rows
        for row_data in rows:
            row_cells = table.add_row().cells
            for i, header in enumerate(headers):
                value = row_data.get(header, "")
                row_cells[i].text = str(value)

    def _add_footer(self, doc: Document) -> None:
        """Add document footer."""
        section = doc.sections[0]
        footer = section.footer
        para = footer.paragraphs[0]
        para.text = (
            f"Generated by M&P Plan Generator V2 | "
            f"Session: {self._session.session_id} | "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def generate_mp_plan_docx(
    session: PlanGeneratorSession,
    output_path: Optional[Path] = None,
    policy_pack: Optional[PolicyPack] = None,
) -> bytes:
    """
    Generate M&P Plan DOCX document.

    Args:
        session: Plan generator session.
        output_path: Optional file path to save document.
        policy_pack: Optional policy pack.

    Returns:
        Document bytes.
    """
    generator = MPPlanDocxGenerator(session, policy_pack)
    return generator.generate(output_path)
