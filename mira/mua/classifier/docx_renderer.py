"""Form 551 DOCX renderer for MUA v10."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .models import Form551Fields

TOP_MARGIN = Cm(2)
BOTTOM_MARGIN = Cm(2)
LEFT_MARGIN = Cm(2.5)
RIGHT_MARGIN = Cm(2.5)

FONT_NAME = "Arial"
BODY_SIZE = Pt(10)
TITLE_SIZE = Pt(14)
HEADER_SIZE = Pt(11)
SECTION_HEADER_COLOR = RGBColor(0x2F, 0x54, 0x96)
FOOTER_SIZE = Pt(8)
FOOTER_COLOR = RGBColor(0x66, 0x66, 0x66)

ISSUE_DISPLAY_NAMES = {
    "FLAM": "Flammability (FLAM) - Promoted Ignition",
    "FS": "Fluid System Compatibility (FS)",
    "TOX": "Offgassing (TOX)",
    "TVS": "Thermal Vacuum Stability (TVS)",
    "SCC": "Stress Corrosion Cracking (SCC)",
    "CORR": "Corrosion (CORR)",
}

RATING_DISPLAY = {
    ("FLAM", "C"): "C (burn length > 6.0 inches per NHB 8060.1C/NASA-STD-6001)",
    (
        "FLAM",
        "I",
    ): "I (fewer than 5 standard samples tested; tested samples at or below 6.0 inches per NHB 8060.1C/NASA-STD-6001)",
    ("FLAM", "U"): "U (unacceptable data per NHB 8060.1C/NASA-STD-6001)",
    ("FS", "U"): "U (untested per NASA-STD-6001B Test 15)",
    ("FS", "A"): "A (passes 48-hour screening)",
    ("TOX", "U"): "U (untested for offgassing per NASA-STD-6001B Test 7)",
    ("SCC", "C"): "C (susceptible under certain conditions)",
    ("SCC", "B"): "B (batch/lot dependent)",
    ("TVS", "C"): "C (fails ASTM E595 criteria)",
    ("CORR", "B"): "B (batch/lot dependent)",
}


def _is_blank(value: str | None) -> bool:
    return value is None or str(value).strip() == ""


def _normalized_value(value: str | None) -> str:
    if _is_blank(value):
        return "[TBD]"
    return str(value).strip()


def _add_run(paragraph, text: str, *, bold: bool = False, italic: bool = False, size=BODY_SIZE, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = FONT_NAME
    run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    return run


def _write_cell(cell, text: str | None, *, label: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    value = _normalized_value(text)
    _add_run(
        paragraph,
        value,
        bold=label,
        italic=(value == "[TBD]"),
    )


def _add_section_header(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    _add_run(
        paragraph,
        text,
        bold=True,
        size=BODY_SIZE,
        color=SECTION_HEADER_COLOR,
    )


def _add_kv_table(doc: Document, rows: int, cols: int):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    return table


def _issue_display(issue: str) -> str:
    return ISSUE_DISPLAY_NAMES.get(issue.upper().strip(), issue)


def _rating_display(issue: str, rating: str) -> str:
    key = (issue.upper().strip(), (rating or "").upper().strip())
    value = RATING_DISPLAY.get(key)
    if value:
        return value
    if _is_blank(rating):
        return "[TBD]"
    return f"{rating}"


def _add_part_numbers(doc: Document, form: Form551Fields) -> None:
    _add_section_header(doc, "PART NUMBER(S)")

    count_text = f"{form.part_count} part occurrence(s):"
    count_paragraph = doc.add_paragraph()
    _add_run(count_paragraph, count_text)

    if form.part_numbers:
        for part in form.part_numbers:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1)
            _add_run(paragraph, f"- {part}")
    else:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1)
        _add_run(paragraph, "- [TBD]", italic=True)


def _add_approvals(doc: Document) -> None:
    _add_section_header(doc, "APPROVALS")

    lines = [
        "Prepared by (Materials Engineer): ____________________  Date: ____________",
        "Reviewed by (Hardware Manager):  ____________________  Date: ____________",
        "Approved by (M&P Branch Chief):  ____________________  Date: ____________",
    ]
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(16)
        _add_run(paragraph, line)


def render_form551_docx(form: Form551Fields) -> bytes:
    """Render a Form551Fields to .docx bytes.

    Parameters
    ----------
    form : Form551Fields
        Populated form fields from the pipeline.

    Returns
    -------
    bytes
        Complete .docx file as bytes. Suitable for StreamingResponse.
    """
    doc = Document()

    section = doc.sections[0]
    section.top_margin = TOP_MARGIN
    section.bottom_margin = BOTTOM_MARGIN
    section.left_margin = LEFT_MARGIN
    section.right_margin = RIGHT_MARGIN

    # Header block
    header_one = doc.add_paragraph()
    header_one.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(header_one, "NATIONAL AERONAUTICS AND SPACE ADMINISTRATION", bold=True, size=HEADER_SIZE)

    header_two = doc.add_paragraph()
    header_two.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(header_two, "MARSHALL SPACE FLIGHT CENTER", bold=True, size=HEADER_SIZE)

    doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, "MATERIAL USAGE AGREEMENT", bold=True, size=TITLE_SIZE)

    # Header metadata table
    meta = _add_kv_table(doc, 3, 4)
    _write_cell(meta.cell(0, 0), "MUA No.", label=True)
    _write_cell(meta.cell(0, 1), form.mua_number)
    _write_cell(meta.cell(0, 2), "Category", label=True)
    _write_cell(meta.cell(0, 3), form.category)

    _write_cell(meta.cell(1, 0), "Date", label=True)
    _write_cell(meta.cell(1, 1), form.date)
    _write_cell(meta.cell(1, 2), "Standard", label=True)
    _write_cell(meta.cell(1, 3), "NASA-STD-6016C")

    _write_cell(meta.cell(2, 0), "Project", label=True)
    _write_cell(meta.cell(2, 1), form.project)
    _write_cell(meta.cell(2, 2), "System", label=True)
    _write_cell(meta.cell(2, 3), form.system)

    _add_section_header(doc, "MATERIAL IDENTIFICATION")
    material_table = _add_kv_table(doc, 6, 2)
    _write_cell(material_table.cell(0, 0), "Material", label=True)
    _write_cell(material_table.cell(0, 1), form.material_designation)
    _write_cell(material_table.cell(1, 0), "MAPTIS Code", label=True)
    _write_cell(material_table.cell(1, 1), form.maptis_code)
    _write_cell(material_table.cell(2, 0), "Specification", label=True)
    _write_cell(material_table.cell(2, 1), form.material_specification)
    _write_cell(material_table.cell(3, 0), "Material Class", label=True)
    _write_cell(material_table.cell(3, 1), form.material_class)
    _write_cell(material_table.cell(4, 0), "Issue", label=True)
    _write_cell(material_table.cell(4, 1), _issue_display(form.issue))
    _write_cell(material_table.cell(5, 0), "Rating", label=True)
    _write_cell(material_table.cell(5, 1), _rating_display(form.issue, form.rating))

    _add_section_header(doc, "OPERATING ENVIRONMENT")
    env_table = _add_kv_table(doc, 2, 4)
    _write_cell(env_table.cell(0, 0), "Media", label=True)
    _write_cell(env_table.cell(0, 1), form.media)
    _write_cell(env_table.cell(0, 2), "Pressure", label=True)
    _write_cell(env_table.cell(0, 3), form.pressure)
    _write_cell(env_table.cell(1, 0), "Temp", label=True)
    _write_cell(env_table.cell(1, 1), form.temperature)
    _write_cell(env_table.cell(1, 2), "Location", label=True)
    _write_cell(env_table.cell(1, 3), form.location)

    _add_part_numbers(doc, form)

    _add_section_header(doc, "APPLICATION")
    application_paragraph = doc.add_paragraph()
    application_value = _normalized_value(form.application)
    _add_run(application_paragraph, application_value, italic=(application_value == "[TBD]"))

    _add_section_header(doc, "RATIONALE")
    rationale_paragraph = doc.add_paragraph()
    rationale_value = _normalized_value(form.rationale)
    _add_run(rationale_paragraph, rationale_value, italic=(rationale_value == "[TBD]"))

    _add_section_header(doc, "DISPOSITION")
    disposition_paragraph = doc.add_paragraph()
    disposition_value = _normalized_value(form.disposition)
    _add_run(disposition_paragraph, disposition_value, italic=(disposition_value == "[TBD]"))

    _add_approvals(doc)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    _add_run(
        footer,
        f"Governing Standard: NASA-STD-6016C | MUA ID: {_normalized_value(form.mua_number)}",
        italic=True,
        size=FOOTER_SIZE,
        color=FOOTER_COLOR,
    )

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()
