"""MIUL -> MUA trigger analysis utilities.

This module identifies non-A MAPTIS ratings in MIUL data and maps each
trigger back to part identity and usage context.
"""

from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

RATING_COLUMNS: tuple[str, ...] = ("FLAM", "SCC", "CORR", "TVS", "TOX", "FS")

# Heuristic risk weights for fast triage bucketing.
# These are intentionally simple and conservative for queueing, not final disposition.
_RISK_WEIGHT_BY_RATING: dict[str, int] = {
    "A": 0,
    "B": 1,
    "C": 2,
    "D": 3,
    "X": 3,
    "U": 2,
    "N": 2,
    "K": 2,
    "I": 2,
}


def _norm_rating(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().upper()


def _row_to_dict(headers: list[str], row_values: list[Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for idx, header in enumerate(headers):
        if header is None:
            continue
        out[str(header)] = row_values[idx] if idx < len(row_values) else None
    return out


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        return float(text)
    except Exception:
        return None


def _trigger_signature(triggers: list[dict[str, str]]) -> str:
    if not triggers:
        return "(none)"
    parts = sorted(f"{t['issue']}:{t['rating']}" for t in triggers)
    return "|".join(parts)


def _risk_score(triggers: list[dict[str, str]]) -> int:
    score = 0
    for trig in triggers:
        score += _RISK_WEIGHT_BY_RATING.get(trig["rating"], 2)
    return score


def _risk_band(score: int) -> str:
    if score >= 5:
        return "high"
    if score >= 2:
        return "medium"
    return "low"


def analyze_miul_workbook(
    workbook_path: str | Path,
    *,
    sheet_name: str = "MIUL",
) -> dict[str, Any]:
    """Analyze a MIUL workbook for non-A MUA triggers.

    Returns a JSON-serializable dict with:
    - workbook metadata
    - rating distributions
    - part-level non-A trigger rows
    - material-code grouped summaries
    """
    path = Path(workbook_path)
    wb = load_workbook(path, data_only=True)
    return _analyze_loaded_workbook(wb, str(path), sheet_name=sheet_name)


def analyze_miul_bytes(
    workbook_bytes: bytes,
    *,
    workbook_name: str = "<upload>",
    sheet_name: str = "MIUL",
) -> dict[str, Any]:
    """Analyze uploaded MIUL workbook bytes for non-A MUA triggers."""
    stream = io.BytesIO(workbook_bytes)
    wb = load_workbook(stream, data_only=True)
    return _analyze_loaded_workbook(wb, workbook_name, sheet_name=sheet_name)


def _analyze_loaded_workbook(
    wb,
    workbook_label: str,
    *,
    sheet_name: str = "MIUL",
) -> dict[str, Any]:
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in workbook. Available: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    header_names = [str(h) for h in headers if h is not None]
    missing = [col for col in RATING_COLUMNS if col not in header_names]
    if missing:
        raise ValueError(
            f"Missing required rating column(s): {missing}. Found headers: {header_names}"
        )

    rating_distributions: dict[str, dict[str, int]] = {
        rc: defaultdict(int) for rc in RATING_COLUMNS
    }
    part_triggers: list[dict[str, Any]] = []

    for row_idx in range(2, ws.max_row + 1):
        row_values = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
        row = _row_to_dict(header_names, row_values)

        non_a: list[dict[str, str]] = []
        for rc in RATING_COLUMNS:
            value = row.get(rc)
            rating = _norm_rating(value)
            bucket = rating if rating else "(blank)"
            rating_distributions[rc][bucket] += 1
            if rating and rating != "A":
                non_a.append({"issue": rc, "rating": rating})

        if not non_a:
            continue

        part_triggers.append(
            {
                "row": row_idx,
                "material_code": row.get("Material Code"),
                "designation": row.get("Designation"),
                "material_specification": row.get("Material Specification"),
                "part_number": row.get("Part Number"),
                "part_description": row.get("Part Description"),
                "nomenclature": row.get("Nomenclature"),
                "nha": row.get("NHA"),
                "environment": row.get("Environment"),
                "max_temperature": row.get("Max Temperature"),
                "min_temperature": row.get("Min Temperature"),
                "max_pressure": row.get("Max Pressure"),
                "min_pressure": row.get("Min Pressure"),
                "fluid_type": row.get("Fluid Type"),
                "system": row.get("System"),
                "subsystem": row.get("Subsystem"),
                "criticality": row.get("Criticality"),
                "overall_evaluation": row.get("Overall Evaluation"),
                "mua_rationale": row.get("MUA Rationale"),
                "remarks": row.get("Remarks"),
                "triggers": non_a,
            }
        )

    # Group by MAPTIS code for MUA batching opportunities
    by_material_code: dict[str, dict[str, Any]] = {}
    for item in part_triggers:
        code = str(item.get("material_code") or "(missing)")
        group = by_material_code.setdefault(
            code,
            {
                "material_code": code,
                "part_count": 0,
                "issue_counts": defaultdict(int),
                "rating_counts": defaultdict(int),
                "part_numbers": [],
                "environments": defaultdict(int),
                "nhas": defaultdict(int),
                "designations": defaultdict(int),
            },
        )
        group["part_count"] += 1
        pn = item.get("part_number")
        if pn is not None:
            group["part_numbers"].append(pn)
        designation = str(item.get("designation") or "").strip()
        if designation:
            group["designations"][designation] += 1
        env = str(item.get("environment") or "")
        if env:
            group["environments"][env] += 1
        nha = str(item.get("nha") or "")
        if nha:
            group["nhas"][nha] += 1
        for trig in item["triggers"]:
            issue = trig["issue"]
            rating = trig["rating"]
            group["issue_counts"][issue] += 1
            group["rating_counts"][f"{issue}:{rating}"] += 1

    # Convert defaultdicts to plain dicts for stable serialization.
    rating_distributions_out = {
        k: dict(sorted(v.items(), key=lambda kv: kv[0]))
        for k, v in rating_distributions.items()
    }
    material_groups_out: list[dict[str, Any]] = []
    for _, group in sorted(
        by_material_code.items(),
        key=lambda kv: (-kv[1]["part_count"], kv[0]),
    ):
        material_groups_out.append(
            {
                "material_code": group["material_code"],
                "part_count": group["part_count"],
                "issue_counts": dict(sorted(group["issue_counts"].items())),
                "rating_counts": dict(sorted(group["rating_counts"].items())),
                "part_numbers": sorted(
                    {str(pn) for pn in group["part_numbers"] if pn is not None}
                ),
                "environments": dict(
                    sorted(group["environments"].items(), key=lambda kv: (-kv[1], kv[0]))
                ),
                "nhas": dict(
                    sorted(group["nhas"].items(), key=lambda kv: (-kv[1], kv[0]))
                ),
                "designations": dict(
                    sorted(group["designations"].items(), key=lambda kv: (-kv[1], kv[0]))
                ),
            }
        )

    return {
        "workbook": workbook_label,
        "sheet": sheet_name,
        "total_parts": ws.max_row - 1,
        "non_a_parts": len(part_triggers),
        "rating_distributions": rating_distributions_out,
        "part_triggers": part_triggers,
        "material_code_groups": material_groups_out,
    }


def build_trigger_buckets(analysis: dict[str, Any]) -> dict[str, Any]:
    """Create triage buckets from analyze_miul_workbook output.

    Buckets are designed for MUA queue triage:
    1) by material code
    2) by trigger signature
    3) by material+signature(+environment)
    """
    part_triggers: list[dict[str, Any]] = analysis.get("part_triggers", [])
    total_parts = int(analysis.get("total_parts", 0))
    non_a_parts = int(analysis.get("non_a_parts", 0))

    by_signature: dict[str, dict[str, Any]] = {}
    by_material_signature: dict[str, dict[str, Any]] = {}

    for part in part_triggers:
        triggers = part.get("triggers", [])
        signature = _trigger_signature(triggers)
        risk_score = _risk_score(triggers)
        risk_band = _risk_band(risk_score)
        material_code = str(part.get("material_code") or "(missing)")
        env = str(part.get("environment") or "")
        bucket_id = f"{material_code}::{signature}::{env}"

        sig_group = by_signature.setdefault(
            signature,
            {
                "trigger_signature": signature,
                "part_count": 0,
                "material_codes": set(),
                "risk_score_max": 0,
                "risk_band": "low",
                "sample_parts": [],
            },
        )
        sig_group["part_count"] += 1
        sig_group["material_codes"].add(material_code)
        sig_group["risk_score_max"] = max(sig_group["risk_score_max"], risk_score)
        sig_group["risk_band"] = _risk_band(sig_group["risk_score_max"])
        if len(sig_group["sample_parts"]) < 5:
            sig_group["sample_parts"].append(
                {
                    "part_number": part.get("part_number"),
                    "part_description": part.get("part_description"),
                    "material_code": material_code,
                }
            )

        ms_group = by_material_signature.setdefault(
            bucket_id,
            {
                "bucket_id": bucket_id,
                "material_code": material_code,
                "trigger_signature": signature,
                "environment": env,
                "part_count": 0,
                "risk_score_max": 0,
                "risk_band": "low",
                "part_numbers": [],
                "nha_values": set(),
                "designations": defaultdict(int),
                "fluid_types": defaultdict(int),
                "temp_min_values": [],
                "temp_max_values": [],
                "pressure_min_values": [],
                "pressure_max_values": [],
            },
        )
        ms_group["part_count"] += 1
        ms_group["risk_score_max"] = max(ms_group["risk_score_max"], risk_score)
        ms_group["risk_band"] = _risk_band(ms_group["risk_score_max"])
        pn = part.get("part_number")
        if pn is not None:
            ms_group["part_numbers"].append(str(pn))
        nha = part.get("nha")
        if nha:
            ms_group["nha_values"].add(str(nha))
        designation = str(part.get("designation") or "").strip()
        if designation:
            ms_group["designations"][designation] += 1
        fluid_type = str(part.get("fluid_type") or "").strip()
        if fluid_type:
            ms_group["fluid_types"][fluid_type] += 1
        tmin = _to_float(part.get("min_temperature"))
        tmax = _to_float(part.get("max_temperature"))
        pmin = _to_float(part.get("min_pressure"))
        pmax = _to_float(part.get("max_pressure"))
        if tmin is not None:
            ms_group["temp_min_values"].append(tmin)
        if tmax is not None:
            ms_group["temp_max_values"].append(tmax)
        if pmin is not None:
            ms_group["pressure_min_values"].append(pmin)
        if pmax is not None:
            ms_group["pressure_max_values"].append(pmax)

    # Material buckets from existing material grouping + risk overlay
    material_buckets: list[dict[str, Any]] = []
    for group in analysis.get("material_code_groups", []):
        material_code = str(group.get("material_code") or "(missing)")
        # aggregate part-level risk for this material code
        scores = [
            _risk_score(p.get("triggers", []))
            for p in part_triggers
            if str(p.get("material_code") or "(missing)") == material_code
        ]
        risk_score_max = max(scores) if scores else 0
        material_buckets.append(
            {
                "material_code": material_code,
                "part_count": int(group.get("part_count", 0)),
                "issue_counts": group.get("issue_counts", {}),
                "rating_counts": group.get("rating_counts", {}),
                "risk_score_max": risk_score_max,
                "risk_band": _risk_band(risk_score_max),
                "environments": group.get("environments", {}),
                "nhas": group.get("nhas", {}),
            }
        )
    material_buckets.sort(
        key=lambda x: (-x["part_count"], -x["risk_score_max"], x["material_code"])
    )

    signature_buckets: list[dict[str, Any]] = []
    for _, g in by_signature.items():
        signature_buckets.append(
            {
                "trigger_signature": g["trigger_signature"],
                "part_count": g["part_count"],
                "material_codes": sorted(g["material_codes"]),
                "material_code_count": len(g["material_codes"]),
                "risk_score_max": g["risk_score_max"],
                "risk_band": g["risk_band"],
                "sample_parts": g["sample_parts"],
            }
        )
    signature_buckets.sort(
        key=lambda x: (-x["part_count"], -x["risk_score_max"], x["trigger_signature"])
    )

    material_signature_buckets: list[dict[str, Any]] = []
    for _, g in by_material_signature.items():
        material_signature_buckets.append(
            {
                "bucket_id": g["bucket_id"],
                "material_code": g["material_code"],
                "trigger_signature": g["trigger_signature"],
                "environment": g["environment"],
                "part_count": g["part_count"],
                "risk_score_max": g["risk_score_max"],
                "risk_band": g["risk_band"],
                "part_numbers": sorted(set(g["part_numbers"])),
                "nha_values": sorted(g["nha_values"]),
                "designations": dict(
                    sorted(g["designations"].items(), key=lambda kv: (-kv[1], kv[0]))
                ),
                "fluid_types": dict(
                    sorted(g["fluid_types"].items(), key=lambda kv: (-kv[1], kv[0]))
                ),
                "temp_min_f": min(g["temp_min_values"]) if g["temp_min_values"] else None,
                "temp_max_f": max(g["temp_max_values"]) if g["temp_max_values"] else None,
                "pressure_min_psia": min(g["pressure_min_values"]) if g["pressure_min_values"] else None,
                "pressure_max_psia": max(g["pressure_max_values"]) if g["pressure_max_values"] else None,
            }
        )
    material_signature_buckets.sort(
        key=lambda x: (-x["part_count"], -x["risk_score_max"], x["bucket_id"])
    )

    return {
        "summary": {
            "total_parts": total_parts,
            "non_a_parts": non_a_parts,
            "no_trigger_parts": max(0, total_parts - non_a_parts),
            "material_bucket_count": len(material_buckets),
            "signature_bucket_count": len(signature_buckets),
            "material_signature_bucket_count": len(material_signature_buckets),
        },
        "material_buckets": material_buckets,
        "signature_buckets": signature_buckets,
        "material_signature_buckets": material_signature_buckets,
    }
