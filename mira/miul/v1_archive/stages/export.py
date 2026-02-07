#!/usr/bin/env python3
"""
Module: Export
Purpose: Stage 5 - Export (JSON → Excel)

BoundState AI 2025
Version: 1.0.0
Classification: BoundState Proprietary

Co-Authored-By: Claude "Beth" <noreply@anthropic.com>
Co-Authored-By: Fran OpenAI 5 <noreply@openai.com>
Co-Authored-By: Grok 4.0 <noreply@x.ai>
Co-Authored-By: Dr. Joseph Wilde, BoundState AI

Compliance: SWE-013, SWE-033, NASA-STD-6016C, MAPTIS Import Spec
Safety Critical: True
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from .merge import MIULDocument, MIULRow

logger = logging.getLogger(__name__)

__all__ = ["export_to_excel", "MAPTIS_COLUMNS"]
__safety_critical__ = True

# MAPTIS 38-column specification per JSC-27301F
MAPTIS_COLUMNS = [
    "Index",
    "Parent",
    "Part Number",
    "Part Description",
    "Nomenclature",
    "NHA",
    "Designation",
    "Material Spec",
    "Process Spec",
    "Form",
    "Material Code",
    "Manufacturer",
    "Qty",
    "Vendor",
    "Vendor PN",
    "Finish",
    "Finish Spec",
    "FLAMM",
    "ODOR",
    "OFF",
    "TOX",
    "COMPAT",
    "SCC",
    "HE",
    "TC-TML",
    "TC-CVCM",
    "TC-WVR",
    "TC-EB",
    "Cert",
    "Test Report",
    "Source Drawing",
    "Find No",
    "Zone",
    "Project",
    "Notes",
    "Date Added",
    "Added By",
    "Revision",
]


def export_to_excel(
    miul: MIULDocument,
    output_path: Optional[Path] = None,
) -> Path:
    """Export MIUL document to MAPTIS-compliant Excel file.

    Stage 5 of the MIUL pipeline. Generates 38-column Excel format
    required for MAPTIS import.

    Formatting matches miul_builder_compliant.py:
    - Dark blue header row
    - Light green MIUL ID rows
    - Light blue assembly rows
    - Plain white for parts/subparts

    Args:
        miul: MIULDocument from Stage 4
        output_path: Where to save Excel file (default: {miul_number}_MIUL.xlsx)

    Returns:
        Path to generated Excel file

    Compliance: MAPTIS Import Spec, JSC-27301F Section 4.3
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl required for Excel export: pip install openpyxl")

    if output_path is None:
        output_path = Path(f"{miul.miul_number}_MIUL.xlsx")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    logger.info(f"[Stage 5] Exporting MIUL to {output_path}")

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIUL"

    # === Styles (match miul_builder_compliant.py format) ===
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Dark blue
    header_font_white = Font(bold=True, size=11, color="FFFFFF")

    miul_id_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")  # Light green
    assembly_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")  # Light blue
    eof_fill = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # === Write Header Row ===
    for col_idx, col_name in enumerate(MAPTIS_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # === Write Data Rows ===
    for row_idx, miul_row in enumerate(miul.rows, 2):
        row_data = _map_row_to_maptis(miul_row, miul)

        for col_idx, col_name in enumerate(MAPTIS_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

        # Apply row-level formatting (match miul_builder_compliant.py)
        row_type = miul_row.row_type

        if row_type == "MIUL_ID":
            for col_idx in range(1, len(MAPTIS_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = miul_id_fill
                ws.cell(row=row_idx, column=col_idx).font = Font(bold=True)

        elif row_type == "ASSEMBLY":
            for col_idx in range(1, len(MAPTIS_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = assembly_fill

        elif row_type == "EOF":
            for col_idx in range(1, len(MAPTIS_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = eof_fill
        # Note: PART and SUBPART rows have no special fill (plain white)

    # === Set Column Widths ===
    column_widths = {
        "Index": 8,
        "Parent": 8,
        "Part Number": 20,
        "Part Description": 35,
        "Nomenclature": 25,
        "NHA": 18,
        "Designation": 20,
        "Material Spec": 18,
        "Process Spec": 18,
        "Form": 12,
        "Material Code": 12,
        "Manufacturer": 15,
        "Qty": 6,
        "Source Drawing": 20,
        "Find No": 10,
        "Project": 12,
        "Notes": 30,
    }

    for col_idx, col_name in enumerate(MAPTIS_COLUMNS, 1):
        width = column_widths.get(col_name, 10)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # === Freeze Header Row ===
    ws.freeze_panes = "A2"

    # === Add Statistics Sheet ===
    ws_stats = wb.create_sheet("Statistics")
    stats = [
        ("MIUL Number", miul.miul_number),
        ("Project", miul.project),
        ("Total Rows", len(miul.rows)),
        ("Assemblies", miul.assembly_count),
        ("Parts", miul.part_count),
        ("Subparts (EPL)", miul.subpart_count),
        ("Generated", miul.generated_at),
        ("Source", miul.hierarchy_source),
        ("Extractor Version", miul.extractor_version),
        ("Export Date", datetime.now().isoformat()),
    ]

    for row_idx, (label, value) in enumerate(stats, 1):
        ws_stats.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws_stats.cell(row=row_idx, column=2, value=value)

    ws_stats.column_dimensions["A"].width = 20
    ws_stats.column_dimensions["B"].width = 40

    # Save
    wb.save(output_path)
    logger.info(f"[Stage 5] Exported {len(miul.rows)} rows to {output_path}")

    return output_path


def _map_row_to_maptis(row: MIULRow, miul: MIULDocument) -> Dict[str, Any]:
    """Map MIULRow fields to MAPTIS column names.

    Args:
        row: Single MIUL row
        miul: Parent document (for project name)

    Returns:
        Dict with MAPTIS column names as keys
    """
    return {
        "Index": row.index,
        "Parent": row.parent if row.parent else "",
        "Part Number": row.part_number,
        "Part Description": row.part_description,
        "Nomenclature": row.nomenclature,
        "NHA": row.nha,
        "Designation": row.designation,
        "Material Spec": row.material_spec,
        "Process Spec": row.process_spec,
        "Form": row.form,
        "Material Code": row.material_code,
        "Manufacturer": row.manufacturer,
        "Qty": row.qty if row.qty and row.qty > 0 else "",
        "Vendor": "",
        "Vendor PN": "",
        "Finish": "",
        "Finish Spec": "",
        "FLAMM": "",
        "ODOR": "",
        "OFF": "",
        "TOX": "",
        "COMPAT": "",
        "SCC": "",
        "HE": "",
        "TC-TML": "",
        "TC-CVCM": "",
        "TC-WVR": "",
        "TC-EB": "",
        "Cert": "",
        "Test Report": "",
        "Source Drawing": row.source_drawing,
        "Find No": row.find_no,
        "Zone": "",
        "Project": miul.project,
        "Notes": "",
        "Date Added": "",
        "Added By": "",
        "Revision": "",
    }
