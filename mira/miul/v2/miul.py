#!/usr/bin/env python3
"""
Module: Miul
Purpose: MIUL Builder v2.2 - Phase 5: MIUL Generation

BoundState AI 2025
Version: 1.0.0
Classification: BoundState Proprietary

Co-Authored-By: Claude "Beth" <noreply@anthropic.com>
Co-Authored-By: Fran OpenAI 5 <noreply@openai.com>
Co-Authored-By: Grok 4.0 <noreply@x.ai>
Co-Authored-By: Dr. Joseph Wilde, BoundState AI

Compliance: SWE-013, NASA-STD-8739.8B
Safety Critical: False
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .models import (
    DrawingGraph,
    DrawingKey,
    DrawingEntity,
    ExplosionPolicy,
    GraphEdge,
    Issue,
    IssueSeverity,
    IssueType,
    MIULEntry,
    NodeType,
    PackageExtract,
)


# =============================================================================
# Configuration
# =============================================================================

# Node types that should appear as MIUL rows (leaf items)
LEAF_NODE_TYPES = {
    NodeType.PIECE_PART,
    NodeType.HARDWARE,
    NodeType.PURCHASED,
    NodeType.STANDARD,
}

# MAPTIS-compatible Excel columns per spec v2.2
MIUL_COLUMNS = [
    "Index",
    "Find No.",
    "Part Number",
    "Description",
    "Qty Per Top",
    "Material",
    "Spec",
    "Finish",
    "Source Drawing",
    "Evidence Chain",
    "Parent",
    "Detail Drawing",
    "Next Assembly",
    "Change Letter",
    "Drawing Source",
    "Material Form",
    "Material Manufacturer",
    "Designation",
    "Material Specification",
    "Process Specification",
    "Environment",
    "Material Code",
    "Standard Part Number",
    "Contractor",
    "System",
    "Subsystem",
    "Max Temperature",
    "Min Temperature",
    "Fluid Type",
    "Associate Contractor",
    "Project",
    "Document Title",
    "Criticality",
    "Line Number",
    "Overall Evaluation",
    "Configuration Test",
    "Max Pressure",
    "Min Pressure",
    "MUA Rationale",
    "Cure Codes",
    "FLAM",
    "SCC",
    "CORR",
    "TVS",
    "TOX",
    "FS",
    "Remarks",
]


# =============================================================================
# Evidence Chain Builder
# =============================================================================

@dataclass
class TraversalContext:
    """Context maintained during tree traversal."""
    path: list[DrawingKey] = field(default_factory=list)
    qty_multiplier: int = 1
    evidence: list[str] = field(default_factory=list)


def build_evidence_chain(
    source_drawing: DrawingKey,
    page_ids: list[str],
    row_index: Optional[int] = None,
) -> list[str]:
    """
    Build an evidence chain for a MIUL entry.

    Format: [drawing_number -> page_id -> row_index]
    """
    chain = [f"drawing:{source_drawing.drawing_number}"]

    if page_ids:
        chain.append(f"page:{page_ids[0]}")

    if row_index is not None:
        chain.append(f"row:{row_index}")

    return chain


# =============================================================================
# Tree Traversal
# =============================================================================

@dataclass
class LeafItem:
    """A leaf item collected during tree traversal."""
    find_no: str
    part_number: str
    description: str
    quantity: int
    material: str
    spec: str
    finish: str
    source_drawing: DrawingKey
    parent_drawing: Optional[DrawingKey]
    evidence_chain: list[str]
    node_type: NodeType
    row_index: Optional[int] = None


def traverse_tree(
    graph: DrawingGraph,
    start_key: DrawingKey,
    context: Optional[TraversalContext] = None,
) -> Iterator[LeafItem]:
    """
    Traverse the drawing tree yielding leaf items.

    Per spec v2.2 per-edge explosion policy:
    - EXPLODE: Recurse into children
    - BLACK_BOX: Count as 1 unit, don't recurse (treat as purchased)
    - PHANTOM: Skip this node, recurse into children

    Yields:
        LeafItem tuples for each leaf part
    """
    if context is None:
        context = TraversalContext()

    context.path.append(start_key)

    entity = graph.nodes.get(start_key)
    if not entity:
        context.path.pop()
        return

    # Get children edges for this drawing
    child_edges = graph.get_children(start_key)

    # Build set of child drawing numbers (subassemblies)
    child_part_numbers = {e.child.drawing_number for e in child_edges}

    # Yield parts that are NOT subassemblies (leaf items)
    for i, part in enumerate(entity.parts_list):
        # Skip if this part is a subassembly that will be handled by recursion
        if part.part_number in child_part_numbers:
            continue

        # Determine node type based on whether part is in graph and has children
        part_key = None
        for key in graph.nodes:
            if key.drawing_number == part.part_number:
                part_key = key
                break

        if part_key and part_key in graph.nodes:
            # Part is a known drawing in our package
            part_children = graph.get_children(part_key)
            if part_children:
                # Has children but no edge from current entity - skip
                # (shouldn't happen normally, but safety check)
                continue
            # Leaf drawing (detail part) - use its entity's node_type or default PIECE_PART
            part_entity = graph.nodes[part_key]
            node_type = part_entity.node_type
            # Override ASSEMBLY default if it's actually a leaf
            if node_type == NodeType.ASSEMBLY and not part_entity.parts_list:
                node_type = NodeType.PIECE_PART
        else:
            # Not in graph - purchased/standard hardware
            part_entity = None
            node_type = NodeType.PURCHASED

        # Determine material: prefer EPL callout, fall back to detail drawing's title block
        material = part.material_callout
        spec = part.spec_callout
        if not material and part_entity:
            # Use detail drawing's own material from title block/notes
            material = part_entity.material_designation
            spec = part_entity.material_spec or spec

        # Build evidence chain
        evidence = build_evidence_chain(
            start_key,
            entity.all_page_ids,
            i,
        )
        evidence.extend(context.evidence)

        # Calculate effective quantity
        effective_qty = part.quantity * context.qty_multiplier

        yield LeafItem(
            find_no=part.find_no or "",
            part_number=part.part_number,
            description=part.description,
            quantity=effective_qty,
            material=material,
            spec=spec,
            finish=part.finish_callout,
            source_drawing=start_key,
            parent_drawing=context.path[-2] if len(context.path) > 1 else None,
            evidence_chain=evidence,
            node_type=node_type,
            row_index=i,
        )

    # Recurse into children based on explosion policy
    for edge in child_edges:
        child_key = edge.child

        if edge.explosion_policy == ExplosionPolicy.BLACK_BOX:
            # Treat as purchased item, don't recurse
            child_entity = graph.nodes.get(child_key)
            description = child_entity.title if child_entity else ""

            evidence = build_evidence_chain(
                start_key,
                entity.all_page_ids,
            )
            evidence.extend(context.evidence)

            effective_qty = edge.qty_per_parent * context.qty_multiplier

            yield LeafItem(
                find_no=edge.find_no or "",
                part_number=child_key.drawing_number,
                description=description,
                quantity=effective_qty,
                material="",
                spec="",
                finish="",
                source_drawing=start_key,
                parent_drawing=start_key,
                evidence_chain=evidence,
                node_type=NodeType.PURCHASED,
            )

        elif edge.explosion_policy == ExplosionPolicy.PHANTOM:
            # Skip this node, recurse into its children
            child_context = TraversalContext(
                path=context.path.copy(),
                qty_multiplier=edge.qty_per_parent * context.qty_multiplier,
                evidence=context.evidence + [f"phantom:{child_key.drawing_number}"],
            )
            yield from traverse_tree(graph, child_key, child_context)

        else:  # EXPLODE (default)
            child_entity = graph.nodes.get(child_key)

            # If child is a leaf node type (PIECE_PART, STANDARD, etc.), yield it directly
            # instead of recursing into an empty parts_list
            if child_entity and child_entity.node_type in LEAF_NODE_TYPES:
                # Find the BOM row for this child to get material/spec/finish
                bom_row = None
                for part in entity.parts_list:
                    if part.part_number == child_key.drawing_number:
                        bom_row = part
                        break

                # Determine material: prefer EPL callout, fall back to detail drawing's title block
                material = bom_row.material_callout if bom_row else ""
                spec = bom_row.spec_callout if bom_row else ""
                if not material and child_entity:
                    # Use detail drawing's own material from title block/notes
                    material = child_entity.material_designation
                    spec = child_entity.material_spec or spec

                evidence = build_evidence_chain(
                    start_key,
                    entity.all_page_ids,
                )
                evidence.extend(context.evidence)

                effective_qty = edge.qty_per_parent * context.qty_multiplier

                yield LeafItem(
                    find_no=edge.find_no or (bom_row.find_no if bom_row else ""),
                    part_number=child_key.drawing_number,
                    description=child_entity.title,
                    quantity=effective_qty,
                    material=material,
                    spec=spec,
                    finish=bom_row.finish_callout if bom_row else "",
                    source_drawing=start_key,
                    parent_drawing=start_key,
                    evidence_chain=evidence,
                    node_type=child_entity.node_type,
                )
            else:
                # Normal recursion for assemblies/subassemblies
                child_context = TraversalContext(
                    path=context.path.copy(),
                    qty_multiplier=edge.qty_per_parent * context.qty_multiplier,
                    evidence=context.evidence + [f"via:{start_key.drawing_number}"],
                )
                yield from traverse_tree(graph, child_key, child_context)

    context.path.pop()


# =============================================================================
# MIUL Generation
# =============================================================================

def generate_miul_entries(
    graph: DrawingGraph,
) -> tuple[list[MIULEntry], list[Issue]]:
    """
    Generate MIUL entries from a DrawingGraph.

    Traverses the tree starting from root, applying per-edge explosion policy.
    Only leaf items (piece parts, hardware, purchased, standards) become rows.
    Subassemblies are NOT counted as parts.

    Returns:
        (list of MIULEntry, list of Issues)
    """
    issues = []
    entries = []
    index = 1

    if not graph.root:
        issues.append(Issue(
            severity=IssueSeverity.WARNING,
            issue_type=IssueType.GRAPH_MISMATCH,
            affected="root",
            message="No root assembly found in graph",
            evidence_page_ids=[],
        ))
        return entries, issues

    # Traverse from root
    for leaf in traverse_tree(graph, graph.root):
        # Check for missing material (per spec: WARNING, include with "MATERIAL TBD")
        material = leaf.material
        if not material and leaf.node_type == NodeType.PIECE_PART:
            material = "MATERIAL TBD"
            issues.append(Issue(
                severity=IssueSeverity.WARNING,
                issue_type=IssueType.LOW_CONFIDENCE,
                affected=leaf.part_number,
                message=f"Missing material for piece part {leaf.part_number}",
                evidence_page_ids=[],
            ))

        entry = MIULEntry(
            index=index,
            find_no=leaf.find_no,
            part_number=leaf.part_number,
            description=leaf.description,
            qty_per_top=leaf.quantity,
            material=material,
            spec=leaf.spec,
            finish=leaf.finish,
            source_drawing=leaf.source_drawing.drawing_number,
            evidence_chain=leaf.evidence_chain,
            parent=leaf.parent_drawing.drawing_number if leaf.parent_drawing else "",
            detail_drawing=leaf.source_drawing.drawing_number,
            next_assembly=leaf.parent_drawing.drawing_number if leaf.parent_drawing else "",
            node_type=leaf.node_type,
        )
        entries.append(entry)
        index += 1

    return entries, issues


def deduplicate_miul_entries(entries: list[MIULEntry]) -> list[MIULEntry]:
    """
    Deduplicate MIUL entries by part number.

    Parts that appear multiple times are consolidated with summed quantities.
    The first occurrence's metadata is preserved.
    """
    seen: dict[str, MIULEntry] = {}

    for entry in entries:
        key = entry.part_number
        if key in seen:
            # Add quantity to existing entry
            seen[key].qty_per_top += entry.qty_per_top
            # Merge evidence chains
            seen[key].evidence_chain.extend(entry.evidence_chain)
        else:
            seen[key] = entry

    # Re-index
    result = list(seen.values())
    for i, entry in enumerate(result, 1):
        entry.index = i

    return result


# =============================================================================
# Excel Export
# =============================================================================

def export_to_excel(
    entries: list[MIULEntry],
    output_path: Path,
    title: str = "MIUL Export",
) -> None:
    """
    Export MIUL entries to Excel file.

    Uses MAPTIS-compatible column format per spec v2.2.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl is required for Excel export")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIUL"

    # Style definitions
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # Write header row
    for col, header in enumerate(MIUL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    # Write data rows
    for row_idx, entry in enumerate(entries, 2):
        row_data = [
            entry.index,
            entry.find_no,
            entry.part_number,
            entry.description,
            entry.qty_per_top,
            entry.material,
            entry.spec,
            entry.finish,
            entry.source_drawing,
            " | ".join(entry.evidence_chain),
            entry.parent,
            entry.detail_drawing,
            entry.next_assembly,
            entry.change_letter,
            entry.drawing_source,
            entry.material_form,
            entry.material_manufacturer,
            entry.designation,
            entry.material_specification,
            entry.process_specification,
            entry.environment,
            entry.material_code,
            entry.standard_part_number,
            "",  # Contractor
            "",  # System
            "",  # Subsystem
            "",  # Max Temperature
            "",  # Min Temperature
            "",  # Fluid Type
            "",  # Associate Contractor
            "",  # Project
            "",  # Document Title
            "",  # Criticality
            "",  # Line Number
            "",  # Overall Evaluation
            "",  # Configuration Test
            "",  # Max Pressure
            "",  # Min Pressure
            "",  # MUA Rationale
            "",  # Cure Codes
            "",  # FLAM
            "",  # SCC
            "",  # CORR
            "",  # TVS
            "",  # TOX
            "",  # FS
            "",  # Remarks
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, ValueError):
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    # Save
    wb.save(output_path)


def export_to_json(
    entries: list[MIULEntry],
    output_path: Path,
) -> None:
    """Export MIUL entries to JSON file."""
    import json

    data = []
    for entry in entries:
        data.append({
            "index": entry.index,
            "find_no": entry.find_no,
            "part_number": entry.part_number,
            "description": entry.description,
            "qty_per_top": entry.qty_per_top,
            "material": entry.material,
            "spec": entry.spec,
            "finish": entry.finish,
            "source_drawing": entry.source_drawing,
            "evidence_chain": entry.evidence_chain,
            "parent": entry.parent,
            "detail_drawing": entry.detail_drawing,
            "next_assembly": entry.next_assembly,
            "node_type": entry.node_type.value,
        })

    with open(output_path, 'w') as f:
        json.dump(data, f, indent=2)


# =============================================================================
# Main MIUL Generation Function
# =============================================================================

def generate_miul(
    package: PackageExtract,
    deduplicate: bool = True,
) -> PackageExtract:
    """
    Generate MIUL for a package.

    This is the main entry point for Phase 5.

    Args:
        package: PackageExtract with graph from Phase 4
        deduplicate: Whether to deduplicate by part number

    Returns:
        Updated PackageExtract with MIUL entries and additional issues
    """
    if not package.graph:
        package.issues.append(Issue(
            severity=IssueSeverity.BLOCKER,
            issue_type=IssueType.GRAPH_MISMATCH,
            affected="graph",
            message="No drawing graph available for MIUL generation",
            evidence_page_ids=[],
        ))
        return package

    # Check for blockers before proceeding
    if package.has_blockers:
        # Don't generate MIUL if there are blockers
        return package

    # Generate MIUL entries
    entries, issues = generate_miul_entries(package.graph)
    package.issues.extend(issues)

    # Deduplicate if requested
    if deduplicate:
        entries = deduplicate_miul_entries(entries)

    package.miul = entries

    return package
