import sys
sys.path.insert(0, r"C:\BoundState")
from docx import Document

doc = Document(r"C:\BoundState\docs\MIRA_v4_BuildSpec_REV1.docx")

# Print all paragraphs
for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)

# Print all tables
for i, table in enumerate(doc.tables):
    print(f"\n=== TABLE {i+1} ===")
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        print(" | ".join(cells))
