#!/usr/bin/env python3
"""Import user-provided FS media mapping data.

The file should contain, at minimum:
  - material code column (e.g. material_code / materialkey)
  - fluid/media column (e.g. fluid_name / media)

Optional:
  - testreportkey
  - testreportnumber
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
import sys

import psycopg2

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from mira.core.db_config import get_db_connection_params


def _norm_key(value: str) -> str:
    return "".join(ch for ch in value.lower().strip() if ch.isalnum())


def _pick(row: dict[str, str], keys: list[str]) -> str:
    for key in keys:
        value = row.get(key, "")
        if value and value.strip().upper() != "NULL":
            return value.strip()
    return ""


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        out: list[dict[str, str]] = []
        for r in reader:
            normalized = {_norm_key(str(k)): ("" if v is None else str(v)) for k, v in r.items()}
            out.append(normalized)
        return out


def _read_xlsx(path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    header_keys = [_norm_key(str(h)) for h in headers]
    out: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        row = {}
        has_value = False
        for col_idx, hk in enumerate(header_keys, start=1):
            value = ws.cell(row=row_idx, column=col_idx).value
            text = "" if value is None else str(value)
            if text.strip():
                has_value = True
            row[hk] = text
        if has_value:
            out.append(row)
    return out


def _load_rows(path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path, sheet_name)
    raise ValueError(f"Unsupported file type: {suffix}. Use .csv or .xlsx.")


def main() -> int:
    parser = argparse.ArgumentParser(description="Import FS media mapping data into public.fs_media_map")
    parser.add_argument("input_file", help="Path to CSV/XLSX with FS media mapping")
    parser.add_argument("--sheet", default=None, help="Worksheet name for XLSX input")
    parser.add_argument("--source", default="user_upload", help="Source tag stored with rows")
    args = parser.parse_args()

    path = Path(args.input_file)
    if not path.exists():
        raise FileNotFoundError(path)

    raw_rows = _load_rows(path, args.sheet)
    parsed: list[tuple[str, str | None, str | None, str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for row in raw_rows:
        material_code = _pick(
            row,
            ["materialcode", "materialkey", "maptiscode", "code"],
        )
        fluid_name = _pick(
            row,
            ["fluidname", "media", "fluid", "testfluid", "medium"],
        )
        if not material_code or not fluid_name:
            continue
        testreportkey = _pick(row, ["testreportkey", "reportkey", "fluctestreportkey"]) or None
        testreportnumber = _pick(row, ["testreportnumber", "reportnumber", "testreportno"]) or None

        dedupe_key = (
            material_code.strip(),
            (testreportkey or "").strip(),
            (testreportnumber or "").strip(),
            fluid_name.strip().upper(),
        )
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        parsed.append(
            (
                material_code.strip(),
                testreportkey.strip() if testreportkey else None,
                testreportnumber.strip() if testreportnumber else None,
                fluid_name.strip(),
                args.source.strip(),
            )
        )

    if not parsed:
        print("No valid rows found. Expected material_code + fluid/media columns.")
        return 1

    cfg = get_db_connection_params(options="-c search_path=maptis,public")
    with psycopg2.connect(**cfg) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS public.fs_media_map (
                    id BIGSERIAL PRIMARY KEY,
                    material_code TEXT NOT NULL,
                    testreportkey TEXT NULL,
                    testreportnumber TEXT NULL,
                    fluid_name TEXT NOT NULL,
                    source TEXT NOT NULL DEFAULT 'user_upload',
                    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            )
            cur.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS fs_media_map_uq
                ON public.fs_media_map (
                    material_code,
                    COALESCE(testreportkey, ''),
                    COALESCE(testreportnumber, ''),
                    UPPER(fluid_name)
                )
                """
            )
            insert_sql = """
                INSERT INTO public.fs_media_map
                    (material_code, testreportkey, testreportnumber, fluid_name, source)
                SELECT %s, %s, %s, %s, %s
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM public.fs_media_map f
                    WHERE f.material_code = %s
                      AND COALESCE(f.testreportkey, '') = COALESCE(%s, '')
                      AND COALESCE(f.testreportnumber, '') = COALESCE(%s, '')
                      AND UPPER(f.fluid_name) = UPPER(%s)
                )
                """
            for material_code, testreportkey, testreportnumber, fluid_name, source in parsed:
                cur.execute(
                    insert_sql,
                    (
                        material_code,
                        testreportkey,
                        testreportnumber,
                        fluid_name,
                        source,
                        material_code,
                        testreportkey,
                        testreportnumber,
                        fluid_name,
                    ),
                )
        conn.commit()

    print(f"Imported {len(parsed)} FS media rows from {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
