"""
Deep structural comparison of Human-verified MIUL vs V3 output.
"""

import openpyxl
from collections import defaultdict, Counter

HUMAN_PATH = 'C:/Users/jwilde2/Desktop/DMLLV/Copy of DMLLV_MIUL_James.xlsx'
V3_PATH = 'C:/Users/jwilde2/Desktop/DMLLV/DMLLV_MIUL_v3.xlsx'


def load_miul(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        if d.get("Part Number") is None and d.get("NHA") is None:
            continue
        rows.append(d)
    return rows


def safe_str(val):
    if val is None:
        return ""
    return str(val).strip()


def print_sep(title, char="=", width=100):
    print()
    print(char * width)
    print("  " + title)
    print(char * width)


def analyze_tree(rows, label):
    print_sep(label + ": TREE STRUCTURE ANALYSIS")
    nha_children = defaultdict(list)
    root_parts = []
    for r in rows:
        nha = safe_str(r.get("NHA"))
        if not nha:
            root_parts.append(r)
        else:
            nha_children[nha].append(r)
    print("\n  Total data rows: %d" % len(rows))
    print("  Root-level rows (NHA is empty): %d" % len(root_parts))
    print("  Child rows (NHA is set): %d" % sum(len(v) for v in nha_children.values()))
    print("  Distinct NHA values: %d" % len(nha_children))
    print("\n  --- Root-level parts ---")
    for r in root_parts:
        part = safe_str(r.get("Part Number"))
        desig = safe_str(r.get("Designation"))
        mat = safe_str(r.get("Material Specification"))
        desc = safe_str(r.get("Part Description"))
        print("    %-30s Desig=%-40s MatSpec=%-30s Desc=%s" % (part, desig, mat, desc))
    print("\n  --- NHA -> Children tree ---")
    for nha in sorted(nha_children.keys()):
        children = nha_children[nha]
        print("\n  NHA=%s (%d children):" % (nha, len(children)))
        for r in children:
            part = safe_str(r.get("Part Number"))
            desig = safe_str(r.get("Designation"))
            mat = safe_str(r.get("Material Specification"))
            print("    %-30s Desig=%-45s MatSpec=%s" % (part, desig, mat))
    return root_parts, nha_children


def analyze_part_frequency(rows, label):
    print_sep(label + ": PART NUMBER FREQUENCY (rows per part)")
    part_counts = Counter()
    part_rows = defaultdict(list)
    for r in rows:
        part = safe_str(r.get("Part Number"))
        if part:
            part_counts[part] += 1
            part_rows[part].append(r)
    multi = {p: c for p, c in part_counts.items() if c > 1}
    print("\n  Total distinct part numbers: %d" % len(part_counts))
    print("  Parts appearing ONCE: %d" % sum(1 for c in part_counts.values() if c == 1))
    print("  Parts appearing MULTIPLE times: %d" % len(multi))
    if multi:
        print("\n  --- Parts with multiple rows ---")
        for part, count in sorted(multi.items(), key=lambda x: -x[1]):
            print("\n    %s (%d rows):" % (part, count))
            for r in part_rows[part]:
                nha = safe_str(r.get("NHA"))
                desig = safe_str(r.get("Designation"))
                mat = safe_str(r.get("Material Specification"))
                print("      NHA=%-20s Desig=%-45s MatSpec=%s" % (nha, desig, mat))
    return part_counts, part_rows


def analyze_dual_row_pattern(rows, label):
    print_sep(label + ": DUAL-ROW PATTERN ANALYSIS")
    part_rows = defaultdict(list)
    for r in rows:
        part = safe_str(r.get("Part Number"))
        if part:
            part_rows[part].append(r)
    dual_count = 0
    print("\n  Looking for parts with DIFFERENT NHA values or different Designation/MatSpec...")
    print("  (dual-row pattern: assembly-level entry + detail-drawing material entry)\n")
    for part, row_list in sorted(part_rows.items()):
        if len(row_list) < 2:
            continue
        nhas = set(safe_str(r.get("NHA")) for r in row_list)
        desigs = [safe_str(r.get("Designation")) for r in row_list]
        matspecs = [safe_str(r.get("Material Specification")) for r in row_list]
        if len(nhas) > 1 or len(set(desigs)) > 1 or len(set(matspecs)) > 1:
            dual_count += 1
            print("  DUAL-ROW: %s" % part)
            for r in row_list:
                nha = safe_str(r.get("NHA"))
                desig = safe_str(r.get("Designation"))
                mat = safe_str(r.get("Material Specification"))
                print("    NHA=%-25s Desig=%-45s MatSpec=%s" % (nha, desig, mat))
    print("\n  Total parts with dual-row: %d" % dual_count)


def analyze_v3_findno_problem(v3_rows):
    print_sep("V3: FIND_NO PROBLEM (MER07543 children)")
    print("\n  V3 has rows under NHA=MER07543-001 where Part Number")
    print("  contains just a number (1, 2, 3...) = find_no, not part number.\n")
    problem = []
    for r in v3_rows:
        nha = safe_str(r.get("NHA"))
        part = safe_str(r.get("Part Number"))
        if nha == "MER07543-001" and part.isdigit():
            problem.append(r)
    if problem:
        print("  Found %d rows with find_no as part number:\n" % len(problem))
        print("  %-15s %-20s %-30s" % ("Part(find_no)", "NHA", "Designation"))
        print("  %s %s %s" % ("-"*15, "-"*20, "-"*30))
        for r in problem:
            print("  %-15s %-20s %-30s" % (safe_str(r.get("Part Number")), safe_str(r.get("NHA")), safe_str(r.get("Designation"))))
    return problem


def analyze_v3_empty_nha(v3_rows):
    print_sep("V3: STANDALONE DETAIL DRAWING ROWS (NHA is empty)")
    print("\n  V3 creates separate rows for detail drawings with NHA=empty.")
    print("  Human puts these as dual-rows under parent NHA instead.\n")
    standalone = []
    for r in v3_rows:
        nha = safe_str(r.get("NHA"))
        desig = safe_str(r.get("Designation"))
        mat = safe_str(r.get("Material Specification"))
        parent = r.get("Parent")
        if not nha and (desig or mat) and parent == 0:
            standalone.append(r)
    print("  Found %d standalone detail drawing rows:\n" % len(standalone))
    print("  %-25s %-45s %s" % ("Part Number", "Designation", "Material Specification"))
    print("  %s %s %s" % ("-"*25, "-"*45, "-"*40))
    for r in standalone:
        print("  %-25s %-45s %s" % (safe_str(r.get("Part Number")), safe_str(r.get("Designation")), safe_str(r.get("Material Specification"))))
    return standalone


def analyze_ocr_errors(human_rows, v3_rows):
    print_sep("OCR ERROR ANALYSIS")
    print("\n" + "-" * 80)
    print("  A. OCR-MANGLED PART NUMBERS IN V3")
    print("-" * 80)
    human_parts = set(safe_str(r.get("Part Number")) for r in human_rows if safe_str(r.get("Part Number")))
    v3_parts = set(safe_str(r.get("Part Number")) for r in v3_rows if safe_str(r.get("Part Number")))
    v3_only = v3_parts - human_parts
    human_only = human_parts - v3_parts
    ocr_map = []
    print("\n  Parts in V3 but NOT in Human (%d):" % len(v3_only))
    for p in sorted(v3_only):
        best = None
        reason = ""
        for hp in human_only:
            if len(p) == len(hp):
                diffs = sum(1 for a, b in zip(p, hp) if a != b)
                if diffs <= 2:
                    best = hp
                    reason = "differs at: %s" % [(i, p[i], hp[i]) for i in range(len(p)) if p[i] != hp[i]]
            elif abs(len(p) - len(hp)) <= 2:
                shorter = p if len(p) < len(hp) else hp
                longer = hp if len(p) < len(hp) else p
                if shorter in longer or longer.startswith(shorter[:max(1, len(shorter)-1)]):
                    best = hp
                    reason = "similar prefix/substring"
        if best:
            ocr_map.append((p, best, reason))
            print("    V3: %-30s -> likely Human: %-30s (%s)" % (p, best, reason))
        else:
            print("    V3: %-30s -> NO MATCH FOUND" % p)
    print("\n  Parts in Human but NOT in V3 (%d):" % len(human_only))
    matched = set(m[1] for m in ocr_map)
    for p in sorted(human_only):
        status = "MATCHED to OCR variant" if p in matched else "MISSING from V3"
        print("    Human: %-30s -> %s" % (p, status))

    print("\n  --- KNOWN OCR PART NUMBER MAPPINGS (manual review) ---")
    known_ocr = [
        ("MS3390-440", "MS9390-440", "OCR: 3 vs 9"),
        ("MS14188N8", "MS14198N6", "OCR: 88->98, 8->6"),
        ("MS19953-6C", "NAS509L6C", "OCR: misread entirely"),
        ("11-885-606", "11-885-06K", "OCR: 606 vs 06K"),
        ("01-885-08N", "01-885-06K", "OCR: 08N vs 06K"),
        ("NAS6706013", "NAS6706U13", "OCR: 0 vs U"),
        ("NAS1351N4L-20", "NAS1351N4-20", "OCR: extra L"),
        ("NAS1352N6-16", "NAS1352N5-12", "OCR: 6->5, 16->12"),
        ("NAS1149E0420R", "NAS1149E0432R", "OCR: 420 vs 432"),
        ("WFM-137-C86", "WSM-137-C86 10.", "OCR: WFM vs WSM"),
        ("M1213302-250", "M12133/2-255", "OCR: missing /, 250->255"),
        ("M121332-255", "M12133/2-255", "OCR: missing /"),
        ("NAS679G113", "NAS6706U13", "OCR: 679G1 vs 6706U"),
        ("WHA-275-588", "WHM-275-S86", "OCR: A->M, 588->S86"),
        ("50854/5CS", "50854SCS", "OCR: /5 vs S"),
        ("MER70570-001", "MER07570-001", "OCR: 70570 vs 07570"),
        ("MER70570-002", "MER07570-???", "OCR: 70570 vs 07570 dash-002"),
        ("MER70575-001", "MER07573-001", "OCR: 70575 vs 07573"),
        ("EH-2-C-65", "EH-22-C86", "OCR: 2-C-65 vs 22-C86"),
        ("MFR07585-001", "MER07585-001", "OCR: MFR vs MER"),
        ("NAS1523H3-16", "NAS1352N5-16", "OCR: 1523H3 vs 1352N5"),
        ("NAS1149E05H6R", "NAS1149E0516R", "OCR: 05H6 vs 0516"),
        ("NAS1523M4-10", "NAS1352N4-10", "OCR: 1523M4 vs 1352N4"),
        ("MT11833-X35", "M12133/2-255", "OCR: MT11833-X35 misread"),
        ("AR293K04", "AS4307K04", "OCR: AR293 vs AS4307"),
        ("F-4006241-04", "F-400521-04", "OCR: 6241 vs 521"),
        ("MSH2090C9-10", "MS21209C4-10", "OCR: MSH2090C9 vs MS21209C4"),
        ("MSH2090C9-15", "MS21209C4/5-15", "OCR: MSH2090C9 vs MS21209C"),
        ("MSH2090C9-20", "MS21209C6-20", "OCR: MSH2090C9 vs MS21209C6"),
        ("MS21902C10-10", "MS21209C4-10", "OCR: 21902C10 vs 21209C4"),
        ("MS21902C4-10", "MS21209C4-10", "OCR: 21902 vs 21209"),
        ("5...", "5.", "OCR: extra dots"),
        ("AS4207K04", "AS4307K04", "OCR: 4207 vs 4307"),
    ]
    for v3_pn, human_pn, reason in known_ocr:
        print("    V3: %-25s -> Human: %-30s (%s)" % (v3_pn, human_pn, reason))
    print("\n  Total known OCR part number errors: %d" % len(known_ocr))

    print("\n" + "-" * 80)
    print("  B. OCR GARBAGE IN MATERIAL SPECIFICATION VALUES")
    print("-" * 80)
    print("\n  V3 has Material Specification values that look like OCR garbage.\n")
    ocr_ms = []
    suspect = ["33982", "33062", "33@B2", "51914", "51588", "56968", "S1914"]
    for r in v3_rows:
        mat = safe_str(r.get("Material Specification"))
        part = safe_str(r.get("Part Number"))
        if not mat:
            continue
        bad = mat in suspect
        if not bad and len(mat) <= 6:
            cleaned = mat.replace("@", "").replace("#", "")
            if cleaned.isdigit():
                bad = True
        if bad:
            ocr_ms.append((part, mat))
    print("  %-25s %-20s %s" % ("Part Number", "OCR MatSpec", "Likely correct (from Human)"))
    print("  %s %s %s" % ("-"*25, "-"*20, "-"*50))
    for part, mat in sorted(set(ocr_ms)):
        human_mat = ""
        for hr in human_rows:
            if safe_str(hr.get("Part Number")) == part:
                hm = safe_str(hr.get("Material Specification"))
                if hm:
                    human_mat = hm
                    break
        print("  %-25s %-20s %s" % (part, mat, human_mat if human_mat else "(not found in human)"))
    print("\n  Total OCR garbage MatSpec entries: %d" % len(ocr_ms))

    all_ms = set()
    for r in v3_rows:
        mat = safe_str(r.get("Material Specification"))
        if mat:
            all_ms.add(mat)
    print("\n  All unique Material Specification values in V3:")
    for ms in sorted(all_ms):
        g = ms in suspect or (len(ms) <= 6 and ms.replace("@", "").replace("#", "").isdigit())
        print("    %s%s" % (ms, " ** OCR GARBAGE **" if g else ""))

    print("\n" + "-" * 80)
    print("  C. DESIGNATION DIFFERENCES")
    print("-" * 80)
    print("\n  Comparing Designation values for parts in both files.\n")
    hd = defaultdict(set)
    vd = defaultdict(set)
    for r in human_rows:
        p = safe_str(r.get("Part Number"))
        d = safe_str(r.get("Designation"))
        if p and d:
            hd[p].add(d)
    for r in v3_rows:
        p = safe_str(r.get("Part Number"))
        d = safe_str(r.get("Designation"))
        if p and d:
            vd[p].add(d)
    common = set(hd.keys()) & set(vd.keys())
    diff_count = 0
    cats = defaultdict(list)
    for part in sorted(common):
        hs = hd[part]
        vs = vd[part]
        if hs != vs:
            diff_count += 1
            h_str = " | ".join(sorted(hs))
            v_str = " | ".join(sorted(vs))
            if any("UNS" in v for v in vs) and not any("UNS" in h for h in hs):
                cat = "MORE_VERBOSE (V3 includes UNS number)"
            elif any("CRES" in v for v in vs) != any("CRES" in h for h in hs):
                cat = "FORMAT_DIFF (CRES vs SST, etc.)"
            elif any("KSI" in v for v in vs) and not any("KSI" in h for h in hs):
                cat = "V3_HAS_KSI (added strength class)"
            elif any("KSI" in h for h in hs) and any("KSI" in v for v in vs):
                cat = "KSI_VALUE_DIFF (different strength)"
            elif vs - hs:
                cat = "DIFFERENT_VALUE"
            else:
                cat = "SUBSET_DIFF"
            cats[cat].append((part, h_str, v_str))
    for cat, items in sorted(cats.items()):
        print("\n  Category: %s (%d parts)" % (cat, len(items)))
        for part, hs, vs in items:
            print("    %s:" % part)
            print("      Human: %s" % hs)
            print("      V3:    %s" % vs)
    print("\n  Total parts with designation differences: %d" % diff_count)


def analyze_missing_standard_parts(human_rows, v3_rows):
    print_sep("MISSING STANDARD PARTS ANALYSIS")
    std_pfx = ["NAS", "MS", "ASME", "M81935", "M12133", "AS4", "NASM", "F-400"]
    v3_ps = set(safe_str(r.get("Part Number")) for r in v3_rows if safe_str(r.get("Part Number")))
    print("\n  Standard parts in Human MIUL MISSING from V3:\n")
    missing = defaultdict(list)
    total = 0
    for r in human_rows:
        part = safe_str(r.get("Part Number"))
        nha = safe_str(r.get("NHA"))
        desig = safe_str(r.get("Designation"))
        mat = safe_str(r.get("Material Specification"))
        is_std = any(part.upper().startswith(p.upper()) for p in std_pfx)
        if part[:4] in ("5085", "9702", "2949"):
            is_std = True
        if part[:3] in ("WHM", "WSM", "EH-"):
            is_std = True
        if is_std and part not in v3_ps:
            missing[nha].append((part, desig, mat))
            total += 1
    for nha in sorted(missing.keys()):
        parts = missing[nha]
        print("  Under NHA=%s (%d missing):" % (nha, len(parts)))
        for part, desig, mat in parts:
            print("    %-30s Desig=%-30s MatSpec=%s" % (part, desig, mat))
        print()
    print("  TOTAL missing standard parts: %d" % total)
    human_ps = set(safe_str(r.get("Part Number")) for r in human_rows if safe_str(r.get("Part Number")))
    print("\n  Standard-looking parts in V3 NOT in Human (potential OCR variants):")
    seen = set()
    for r in v3_rows:
        part = safe_str(r.get("Part Number"))
        is_std = any(part.upper().startswith(p.upper()) for p in std_pfx)
        if is_std and part and part not in human_ps and part not in seen:
            seen.add(part)
            print("    V3: %-30s NHA=%-20s Desig=%s" % (part, safe_str(r.get("NHA")), safe_str(r.get("Designation"))))


def analyze_v3_structural(v3_rows, human_rows):
    print_sep("V3 STRUCTURAL ISSUES SUMMARY")
    h_root = any(safe_str(r.get("Part Number")) == "MER07540" and not safe_str(r.get("NHA")) for r in human_rows)
    v_root = any(safe_str(r.get("Part Number")) == "MER07540" and not safe_str(r.get("NHA")) for r in v3_rows)
    print("\n  1. Top-level MER07540 assembly row:")
    print("     Human has it: %s" % h_root)
    print("     V3 has it:    %s" % v_root)
    if not v_root:
        print("     -> V3 starts at MER07540-001, missing top-level wrapper")
    v3_icd = [r for r in v3_rows if safe_str(r.get("Part Number")) == "MER07540-ICD"]
    if v3_icd:
        print("\n  2. V3 has MER07540-ICD row (not in Human) - ICD ref, not MIUL part")
    v3_no = [r for r in v3_rows if not safe_str(r.get("NHA"))]
    h_no = [r for r in human_rows if not safe_str(r.get("NHA"))]
    print("\n  3. Rows with empty NHA:")
    print("     Human: %d (top-level only)" % len(h_no))
    print("     V3:    %d (includes standalone detail-drawing rows)" % len(v3_no))
    h543 = [r for r in human_rows if safe_str(r.get("NHA")) == "MER07543-001"]
    v543 = [r for r in v3_rows if safe_str(r.get("NHA")) == "MER07543-001"]
    v543f = [r for r in v543 if safe_str(r.get("Part Number")).isdigit()]
    print("\n  4. Children of MER07543-001:")
    print("     Human: %d children (real part numbers)" % len(h543))
    print("     V3:    %d children (%d with find_no as part number)" % (len(v543), len(v543f)))
    print("\n  5. Parts under wrong NHA in V3:")
    hm = defaultdict(set)
    vm = defaultdict(set)
    for r in human_rows:
        p = safe_str(r.get("Part Number"))
        n = safe_str(r.get("NHA"))
        if p and n:
            hm[p].add(n)
    for r in v3_rows:
        p = safe_str(r.get("Part Number"))
        n = safe_str(r.get("NHA"))
        if p and n:
            vm[p].add(n)
    for part in sorted(set(hm.keys()) & set(vm.keys())):
        extra = vm[part] - hm[part]
        if extra:
            print("     %s: V3 NHA=%s not in Human (Human NHA=%s)" % (part, extra, hm[part]))


def print_key_findings(human_rows, v3_rows):
    print_sep("KEY FINDINGS: PRIORITIZED ISSUES TO FIX IN V3 BUILDER", char="*")
    findings = [
        ("CRITICAL", "Find Number used as Part Number for MER07543 children",
         "V3 uses find_no (1,2,3...) instead of actual part numbers for all 29 children of MER07543-001. Builder must resolve find numbers to part numbers using EPL data."),
        ("CRITICAL", "Widespread OCR errors in part numbers (~33 known errors)",
         "At least 33 part numbers OCR-mangled: MS3390-440->MS9390-440, MER70570-001->MER07570-001, NAS6706013->NAS6706U13, MSH2090C9-10->MS21209C4-10."),
        ("CRITICAL", "OCR garbage in Material Specification column",
         "Multiple rows have nonsensical MatSpec: 33982, 33062, 33@B2, 51914, 51588, 56968. OCR artifacts. Need spec pattern validation (AMS, ASTM, SAE)."),
        ("HIGH", "Missing top-level assembly hierarchy (MER07540 -> MER07540-001)",
         "Human has MER07540 as root with MER07540-001 as child. V3 starts at MER07540-001. Builder needs top-level assembly row."),
        ("HIGH", "Standalone detail-drawing rows instead of dual-row pattern",
         "V3 creates ~25 standalone rows with NHA=empty for detail drawings. Human puts these as SECOND rows under parent NHA with detailed material specs."),
        ("HIGH", "Missing standard parts from EPL tables",
         "Multiple standard parts missing: NAS509L6C, NAS1149E0432R, NAS1352N5-12, NAS1351N4H16, NAS1352N4-12/14, NAS509-4C, NASM20995N32, McMaster parts, WHM-275-S86, WHM-244-S86."),
        ("HIGH", "Spurious MER07540-ICD row in V3",
         "V3 includes MER07540-ICD not in Human MIUL. ICD reference, not a MIUL part."),
        ("MEDIUM", "Designation values differ in format/verbosity",
         "V3 designations differ: UNS numbers included, wrong KSI strength (180 vs 150/160), wrong materials (INCONEL 718 vs 625 for MER07574-001), VESPEL SP-21 vs SP-3."),
        ("MEDIUM", "Parts misplaced in assembly tree (wrong NHA)",
         "MER07580-001 under MER07560-001 in V3 (7075 AL) but under MER07573-001 in Human (INCONEL 625). MER07582/83-001 also wrong NHA."),
        ("MEDIUM", "VESPEL SP-21 vs SP-3 material designation error",
         "V3 uses VESPEL SP-21 for MER07578/79-001, Human uses VESPEL SP-3 POLYIMIDE RESIN. Different grades - material error."),
        ("LOW", "MER07576-001 vs MER07576 dash number inconsistency",
         "V3 uses MER07576-001, Human uses MER07576 (no dash). Affects child NHA refs."),
        ("LOW", "Duplicate rows for NAS509-6C in V3",
         "V3 has NAS509-6C twice under MER07540-001. De-duplication needed."),
    ]
    for i, (sev, title, desc) in enumerate(findings, 1):
        print("\n  %d. [%s] %s" % (i, sev, title))
        print("     %s" % desc)
    crit = sum(1 for f in findings if f[0] == "CRITICAL")
    high = sum(1 for f in findings if f[0] == "HIGH")
    med = sum(1 for f in findings if f[0] == "MEDIUM")
    low = sum(1 for f in findings if f[0] == "LOW")
    print("\n" + "*" * 100)
    print("  SUMMARY: %d CRITICAL, %d HIGH, %d MEDIUM, %d LOW issues" % (crit, high, med, low))
    print("*" * 100)


def main():
    print("=" * 100)
    print("  DEEP STRUCTURAL COMPARISON: Human MIUL vs V3 Output")
    print("  Human: Copy of DMLLV_MIUL_James.xlsx")
    print("  V3:    DMLLV_MIUL_v3.xlsx")
    print("=" * 100)
    human_rows = load_miul(HUMAN_PATH)
    v3_rows = load_miul(V3_PATH)
    print("\n  Loaded Human MIUL: %d data rows" % len(human_rows))
    print("  Loaded V3 MIUL:    %d data rows" % len(v3_rows))
    analyze_tree(human_rows, "HUMAN")
    analyze_tree(v3_rows, "V3")
    analyze_part_frequency(human_rows, "HUMAN")
    analyze_part_frequency(v3_rows, "V3")
    analyze_dual_row_pattern(human_rows, "HUMAN")
    analyze_dual_row_pattern(v3_rows, "V3")
    analyze_v3_findno_problem(v3_rows)
    analyze_v3_empty_nha(v3_rows)
    analyze_ocr_errors(human_rows, v3_rows)
    analyze_missing_standard_parts(human_rows, v3_rows)
    analyze_v3_structural(v3_rows, human_rows)
    print_key_findings(human_rows, v3_rows)


if __name__ == "__main__":
    main()
